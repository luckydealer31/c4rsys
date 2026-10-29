"""Reports, parts О1–О2: the area of data, period, export and the «Company overview» (ТЗ § 21.1,
§ 21.2.6–21.2.11, § 21.3.1, § 23.8.5; customer's decisions О-1, О-2, О-5)."""

import io
from typing import Any

import pytest
from django.test import Client
from django.utils import timezone
from openpyxl import load_workbook

from tests.accounts.helpers import post
from tests.companies.helpers import hire, key
from tests.orders.test_close import _pay, desk, frozen_clock  # noqa: F401
from tests.rentals.scenario import Scenario, add_location, issue

pytestmark = pytest.mark.django_db


def _report(client: Client, s: Scenario, query: str = "") -> Any:
    return client.get(f"/api/c/{s.cmp}/reports/overview?{query}")


def _blocks(client: Client, s: Scenario, query: str = "") -> dict[str, dict[str, Any]]:
    response = _report(client, s, query)
    assert response.status_code == 200, response.content
    return {b["code"]: b for b in response.json()["blocks"]}


def test_overview_counts_and_whole_area(owner: Client, s: Scenario, desk: str) -> None:  # noqa: F811
    rental = issue(owner, s).json()["rental"]
    paid = _pay(owner, s, desk, rental["order_id"], "500")
    assert paid.status_code == 201, paid.content
    blocks = _blocks(owner, s)
    assert blocks["rental.active"]["value"] == 1
    assert blocks["fleet.in_rental"]["value"] == 1 and blocks["fleet.total"]["value"] == 1
    assert blocks["fleet.available"]["value"] == 0
    assert blocks["rental.active"]["area"]["whole"] is True
    assert blocks["finance.receipts"]["rows"] == [{"currency": "THB", "amount": "500.00"}]
    assert blocks["finance.balances"]["rows"][0]["balance"] == "500.00"
    # The paid part is ahead of the charges of an active rental: no debt yet, a counted zero.
    assert blocks["finance.debt"]["value"] == 0
    # Another period: a counted zero, not «no data» (§ 21.2.6).
    later = _blocks(owner, s, "start=2030-01-01&end=2030-01-31")
    assert later["finance.receipts"]["value"] == 0 and later["finance.receipts"]["rows"] == []


def test_area_follows_the_rights(owner: Client, s: Scenario) -> None:
    other = add_location(owner, s.cmp, "Kata")
    issue(owner, s)
    reader, _ = hire(
        owner, s.cmp, "r@example.com", [{"type": "grant", "permission": "reports.view"}]
    )
    # Without the rights of the sections every block is hidden, not zero (§ 21.1.1).
    assert _blocks(reader, s) == {}
    local, _ = hire(
        owner,
        s.cmp,
        "k@example.com",
        [
            {"type": "grant", "permission": "reports.view"},
            {"type": "grant", "permission": "rental.view", "locations": [other]},
        ],
    )
    blocks = _blocks(local, s)
    assert blocks["rental.active"]["value"] == 0
    assert blocks["rental.active"]["area"] == {"whole": False, "locations": ["Kata"]}
    assert "finance.receipts" not in blocks
    # A location out of reach is not shown.
    assert _blocks(local, s, f"location={s.location}") == {}


def test_rights_period_and_export(owner: Client, s: Scenario) -> None:
    viewer, _ = hire(
        owner, s.cmp, "v@example.com", [{"type": "grant", "permission": "rental.view"}]
    )
    assert _report(viewer, s).status_code == 403
    assert (
        _report(owner, s, "start=2026-02-01&end=2026-01-01").json()["error"]["code"]
        == "PERIOD_INVALID"
    )
    data = _report(owner, s, "tz=Europe/Moscow").json()
    assert data["filters"]["zone"] == "Europe/Moscow"
    today = timezone.now().date().replace(day=1).isoformat()
    assert _report(owner, s).json()["filters"]["start"][:8] == today[:8]

    reader, _ = hire(
        owner, s.cmp, "r@example.com", [{"type": "grant", "permission": "reports.view"}]
    )
    refused = post(reader, f"/api/c/{s.cmp}/reports/overview/xlsx", {}, key())
    assert refused.status_code == 403
    labels = {"overview": "Обзор компании", "rental.active": "Активные аренды", "period": "Период"}
    response = post(owner, f"/api/c/{s.cmp}/reports/overview/xlsx", {"labels": labels}, key())
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    values = [row[0] for row in sheet.iter_rows(values_only=True)]
    assert values[0] == "Обзор компании" and values[1] == "Период" and "Активные аренды" in values
