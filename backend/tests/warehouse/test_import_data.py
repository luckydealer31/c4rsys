"""XLSX import of vehicles, clients and counterparties, corrected files, cancel and the keeping
of source files (ТЗ § 22.2–22.5; block 6, part К2)."""

import io
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook

from c4rsys.fleet.models import OdometerReading, Vehicle, VehicleCard
from c4rsys.parties.models import ClientCard, Counterparty
from c4rsys.warehouse import imports
from c4rsys.warehouse.models import ImportOperation
from tests.accounts.helpers import post
from tests.companies.helpers import hire, key
from tests.rentals.scenario import add_client, build

pytestmark = pytest.mark.django_db


def workbook(content: bytes, sheet_name: str, rows: list[dict[str, str]]) -> bytes:
    book = load_workbook(io.BytesIO(content))
    sheet = book[sheet_name]
    header = [c.value for c in sheet[1]]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append([row.get(str(code), "") for code in header])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def template(client: Client, cmp: str, kind: str) -> bytes:
    response = client.get(f"/api/c/{cmp}/imports/template/{kind}")
    assert response.status_code == 200, response.content
    return response.content


def upload(client: Client, cmp: str, kind: str, content: bytes) -> Any:
    return client.post(
        f"/api/c/{cmp}/imports",
        data={"kind": kind, "file": SimpleUploadedFile(f"{kind}.xlsx", content)},
        headers=key(),
    )


def checked(client: Client, cmp: str, kind: str, rows: list[dict[str, str]]) -> dict[str, Any]:
    content = workbook(template(client, cmp, kind), kind, rows)
    response = upload(client, cmp, kind, content)
    assert response.status_code == 201, response.content
    body: dict[str, Any] = response.json()["import"]
    return body


def run(client: Client, cmp: str, operation: dict[str, Any], step: str, **data: Any) -> Any:
    return post(
        client,
        f"/api/c/{cmp}/imports/{operation['id']}/{step}",
        {"expected_version": operation["version"], **data},
        key(),
    )


def test_vehicles_are_checked_matched_and_created(owner: Client) -> None:
    s = build(owner)
    book = load_workbook(io.BytesIO(template(owner, s.cmp, "vehicles")))
    assert book["locations"]["A2"].value == s.location  # the IDs of the company's locations
    existing = post(
        owner,
        f"/api/c/{s.cmp}/vehicles",
        {
            "name": "Old",
            "location_id": s.location,
            "odometer_unit": "km",
            "vin": "JTDBR32E720012345",
        },
        key(),
    ).json()["vehicle"]
    base = {"location_id": s.location, "relation": "own", "odometer_unit": "km"}
    operation = checked(
        owner,
        s.cmp,
        "vehicles",
        [
            {**base, "make": "Toyota", "model": "Camry", "plate": "1AB 234", "odometer": "35000"},
            {**base, "vin": "JTDBR32E720012345"},  # the same VIN as a card of the company
            {**base, "relation": "", "name": "No relation"},
            {**base, "location_id": "00000000-0000-0000-0000-000000000000", "name": "Elsewhere"},
            {
                **base,
                "relation": "managed",
                "name": "Owner car",
                "owner_kind": "person",
                "owner_name": "Somchai",
            },
            {**base},  # no name and nothing to form it from
        ],
    )
    rows = operation["rows"]
    assert [r["action"] for r in rows] == ["create", "skip", "error", "error", "create", "error"]
    assert rows[1]["matches"] == [{"id": existing["id"], "name": "Old", "by": "vin"}]
    assert rows[5]["errors"] == [{"field": "name", "code": "NAME_NOT_DETERMINED"}]

    decided = run(
        owner,
        s.cmp,
        operation,
        "decisions",
        decisions=[{"row": "R0002", "action": "link", "target": existing["id"]}],
    ).json()
    posted = run(owner, s.cmp, decided["import"], "post")
    assert posted.status_code == 200, posted.content
    body = posted.json()["import"]
    assert (body["loaded_count"], body["skipped_count"]) == (3, 3)
    camry = VehicleCard.objects.get(name="Toyota Camry 1AB 234")
    assert camry.vehicle.make_id and camry.accounting_location_id is not None
    assert OdometerReading.objects.get(vehicle_card=camry).value == 35000
    assert VehicleCard.objects.get(name="Owner car").relation == "managed"
    assert body["rows"][1]["result"] == "linked"


def test_vehicle_id_and_vin_pointing_to_different_cards_is_a_conflict(owner: Client) -> None:
    s = build(owner)
    first = post(
        owner,
        f"/api/c/{s.cmp}/vehicles",
        {"name": "A", "location_id": s.location, "odometer_unit": "km"},
        key(),
    ).json()["vehicle"]
    post(
        owner,
        f"/api/c/{s.cmp}/vehicles",
        {"name": "B", "location_id": s.location, "odometer_unit": "km", "vin": "JTDBR32E720099999"},
        key(),
    )
    vhl = VehicleCard.objects.get(pk=first["id"]).vehicle.public_id
    operation = checked(
        owner,
        s.cmp,
        "vehicles",
        [
            {
                "vehicle_id": vhl,
                "vin": "JTDBR32E720099999",
                "location_id": s.location,
                "relation": "own",
                "odometer_unit": "km",
                "name": "C",
            }
        ],
    )
    assert operation["rows"][0]["conflict"] and operation["rows"][0]["action"] == "skip"
    refused = run(
        owner,
        s.cmp,
        operation,
        "decisions",
        decisions=[{"row": "R0001", "action": "create_duplicate"}],
    )
    assert refused.json()["error"]["code"] == "IMPORT_ROW_ID_CONFLICT"


def test_clients_match_by_contacts_not_by_name(owner: Client) -> None:
    s = build(owner)
    existing = add_client(owner, s.cmp, "Ivan", "+66899990000")
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [
            {
                "kind": "person",
                "first_name": "Anna",
                "last_name": "Lee",
                "phone": "+66 89 999 0000",
            },
            {"kind": "person", "first_name": "Ivan", "last_name": "Petrov", "email": "i@ex.com"},
            {
                "kind": "company",
                "company_name": "Sun Tours",
                "email": "sun@ex.com",
                "address": "Patong",
                "external_id": "CL-7",
                "c4rsys_id": "USR-NOPE",
            },
            {"kind": "person", "first_name": "No", "last_name": "Contact"},
        ],
    )
    rows = operation["rows"]
    # The same phone written differently is a match; the same name alone is not (§ 22.2.3).
    assert rows[0]["matches"][0] == {"id": existing, "name": "Petrov Ivan", "by": "phone"}
    assert rows[1]["matches"] == []
    assert rows[3]["errors"] == [{"field": "phone", "code": "CONTACT_REQUIRED"}]
    decided = run(
        owner,
        s.cmp,
        operation,
        "decisions",
        decisions=[{"row": "R0001", "action": "create_duplicate"}],
    ).json()["import"]
    body = run(owner, s.cmp, decided, "post").json()["import"]
    assert (body["loaded_count"], body["skipped_count"]) == (3, 1)
    assert ClientCard.objects.get(first_name="Anna").possible_duplicate
    sun = ClientCard.objects.get(company_name="Sun Tours")
    assert (sun.address, sun.external_id, sun.linked_user_id) == ("Patong", "CL-7", None)

    # The external ID finds the card again; another card's phone in the same row is a conflict.
    again = checked(
        owner,
        s.cmp,
        "clients",
        [
            {
                "kind": "company",
                "company_name": "Sun",
                "external_id": "CL-7",
                "phone": "+66899990000",
            }
        ],
    )
    assert again["rows"][0]["conflict"]
    assert {m["by"] for m in again["rows"][0]["matches"]} == {"external_id", "phone"}


def test_counterparties_are_searched_in_their_own_list(owner: Client) -> None:
    s = build(owner)
    add_client(owner, s.cmp, "Ivan", "+66899990000")
    operation = checked(
        owner,
        s.cmp,
        "counterparties",
        [
            {
                "kind": "company",
                "company_name": "Parts Co",
                "phone": "+66812345678",
                "phone_messenger": "Telegram",
                "registration": "TAX 123",
                "category": "Parts",
            }
        ],
    )
    assert operation["rows"][0]["matches"] == []  # a client with the phone does not count
    body = run(owner, s.cmp, operation, "post").json()["import"]
    assert body["rows"][0]["result"] == "created"
    card = Counterparty.objects.get(company_name="Parts Co")
    assert (card.registration, card.category) == ("TAX 123", "Parts")
    # A messenger is bound to the phone number (decision of 30.10.2026).
    assert card.contacts.get().messenger == "telegram"


def test_messenger_belongs_to_a_phone(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [
            {
                "kind": "person",
                "first_name": "A",
                "last_name": "B",
                "email": "a@ex.com",
                "phone_messenger": "whatsapp",
            },
            {
                "kind": "person",
                "first_name": "C",
                "last_name": "D",
                "phone": "+6681111",
                "phone_messenger": "pigeon",
            },
        ],
    )
    assert operation["rows"][0]["errors"] == [
        {"field": "phone_messenger", "code": "MESSENGER_WITHOUT_PHONE"}
    ]
    assert operation["rows"][1]["errors"] == [{"field": "phone", "code": "MESSENGER_INVALID"}]
    created = post(
        owner,
        f"/api/c/{s.cmp}/clients",
        {
            "kind": "person",
            "first_name": "E",
            "last_name": "F",
            "contacts": [
                {"kind": "phone", "value": "+6682222", "messenger": "whatsapp"},
                {"kind": "email", "value": "e@ex.com"},
            ],
        },
        key(),
    )
    assert created.status_code == 201, created.content
    contacts = created.json()["client"]["contacts"]
    assert {"kind": "phone", "value": "+6682222", "messenger": "whatsapp"} in contacts
    refused = post(
        owner,
        f"/api/c/{s.cmp}/clients",
        {
            "kind": "person",
            "first_name": "G",
            "last_name": "H",
            "contacts": [{"kind": "email", "value": "g@ex.com", "messenger": "telegram"}],
        },
        key(),
    )
    assert refused.json()["error"]["code"] == "MESSENGER_INVALID"


def test_body_chassis_and_frame_numbers(owner: Client) -> None:
    s = build(owner)
    created = post(
        owner,
        f"/api/c/{s.cmp}/vehicles",
        {
            "name": "Truck",
            "location_id": s.location,
            "odometer_unit": "km",
            "body_number": " ab123 ",
            "frame_number": "FR-9",
        },
        key(),
    )
    assert created.status_code == 201, created.content
    card = VehicleCard.objects.get(pk=created.json()["vehicle"]["id"])
    assert (card.vehicle.body_number, card.vehicle.frame_number) == ("AB123", "FR-9")
    base = {"location_id": s.location, "relation": "own", "odometer_unit": "km"}
    operation = checked(
        owner,
        s.cmp,
        "vehicles",
        [
            {**base, "name": "Same body", "body_number": "ab123"},
            {**base, "name": "Body as chassis", "chassis_number": "AB123"},  # another kind
        ],
    )
    assert operation["rows"][0]["matches"][0]["by"] == "body_number"
    assert operation["rows"][1]["matches"] == []
    body = run(owner, s.cmp, operation, "post").json()["import"]
    assert body["rows"][1]["result"] == "created"
    other = VehicleCard.objects.get(name="Body as chassis")
    assert other.vehicle.chassis_number == "AB123"


def test_corrected_file_is_matched_by_row_codes(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [
            {"kind": "person", "first_name": "Anna", "last_name": "Lee", "phone": "+6681111"},
            {"kind": "person", "first_name": "Bob", "last_name": "", "phone": "+6682222"},
            {"kind": "person", "first_name": "Cat", "last_name": "", "phone": "+6683333"},
        ],
    )
    posted = run(owner, s.cmp, operation, "post").json()["import"]
    assert posted["loaded_count"] == 1
    url = f"/api/c/{s.cmp}/imports/{operation['id']}"
    correction = owner.get(f"{url}/correction")
    assert correction.status_code == 200
    book = load_workbook(io.BytesIO(correction.content))
    sheet = book["clients"]
    assert [c.value for c in sheet[1]][:2] == ["row_code", "kind"]
    assert [sheet.cell(row=r, column=1).value for r in (2, 3)] == ["R0002", "R0003"]
    assert book["problems"]["E2"].value == "FIELD_REQUIRED"
    # Fixed Bob; Cat stays wrong; a row of Anna (already loaded) and a row without a code.
    header = [c.value for c in sheet[1]]
    sheet.delete_rows(2, 2)
    for values in (
        {
            "row_code": "R0003",
            "kind": "person",
            "first_name": "Cat",
            "last_name": "",
            "phone": "+6683333",
        },
        {
            "row_code": "R0002",
            "kind": "person",
            "first_name": "Bob",
            "last_name": "Ray",
            "phone": "+6682222",
        },
        {
            "row_code": "R0001",
            "kind": "person",
            "first_name": "Anna",
            "last_name": "Changed",
            "phone": "+6681111",
        },
        {
            "row_code": "",
            "kind": "person",
            "first_name": "Dan",
            "last_name": "Fox",
            "phone": "+6684444",
        },
    ):
        sheet.append([values.get(str(c), "") for c in header])
    output = io.BytesIO()
    book.save(output)
    reuploaded = owner.post(
        f"{url}/reupload",
        data={
            "expected_version": str(posted["version"]),
            "file": SimpleUploadedFile("fixed.xlsx", output.getvalue()),
        },
        headers=key(),
    )
    assert reuploaded.status_code == 200, reuploaded.content
    body = reuploaded.json()["import"]
    codes = {w["code"] for w in body["warnings"]}
    assert {"ROW_ALREADY_POSTED", "ROW_CODE_MISSING", "CORRECTED_FILE"} <= codes
    assert [r["action"] for r in body["rows"]] == ["create", "create", "error"]
    final = run(owner, s.cmp, body, "post").json()["import"]
    assert (final["loaded_count"], final["skipped_count"]) == (2, 1)
    assert ClientCard.objects.get(first_name="Anna").last_name == "Lee"  # not overwritten
    assert ClientCard.objects.filter(first_name="Bob", last_name="Ray").exists()
    assert not ClientCard.objects.filter(first_name="Dan").exists()


def test_cancel_is_whole_and_stops_at_later_history(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [{"kind": "person", "first_name": "Anna", "last_name": "Lee", "phone": "+6681111"}],
    )
    posted = run(owner, s.cmp, operation, "post").json()["import"]
    card = ClientCard.objects.get(first_name="Anna")
    changed = owner.patch(
        f"/api/c/{s.cmp}/clients/{card.id}",
        data={"last_name": "Ray", "expected_version": 1},
        content_type="application/json",
        headers=key(),
    )
    assert changed.status_code == 200, changed.content
    refused = run(owner, s.cmp, posted, "cancel")
    assert refused.json()["error"]["code"] == "IMPORT_CANCEL_BLOCKED"
    assert refused.json()["error"]["params"]["blockers"][0]["action"] == "client.changed"

    vehicles = checked(
        owner,
        s.cmp,
        "vehicles",
        [
            {
                "location_id": s.location,
                "relation": "own",
                "odometer_unit": "km",
                "name": "Imported",
                "plate": "9ZZ 999",
                "odometer": "100",
            }
        ],
    )
    done = run(owner, s.cmp, vehicles, "post").json()["import"]
    cancelled = run(owner, s.cmp, done, "cancel")
    assert cancelled.json()["import"]["state"] == "cancelled", cancelled.content
    assert not VehicleCard.objects.filter(name="Imported").exists()
    assert not Vehicle.objects.filter(plates__number="9ZZ 999").exists()


def test_files_are_removed_six_months_after_the_first_upload(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [{"kind": "person", "first_name": "Anna", "last_name": "Lee", "phone": "+6681111"}],
    )
    url = f"/api/c/{s.cmp}/imports/{operation['id']}/file"
    assert owner.get(url).status_code == 200
    created = ImportOperation.objects.get(pk=operation["id"]).created_at
    assert imports.remove_old_files(created + timedelta(days=150)) == 0
    assert imports.remove_old_files(imports.shift_months(created, 6)) == 1
    assert owner.get(url).json()["error"]["code"] == "FILE_REMOVED"
    detail = owner.get(f"/api/c/{s.cmp}/imports/{operation['id']}").json()["import"]
    assert detail["files_removed"] and len(detail["rows"]) == 1
    assert imports.shift_months(datetime(2026, 8, 31, tzinfo=UTC), 6).day == 28


def test_import_needs_the_right_of_the_section(owner: Client) -> None:
    s = build(owner)
    rights = [
        {
            "type": "grant",
            "permission": p,
            "locations": "all",
            "contexts": "all",
            "cash_accounts": "all",
        }
        for p in ("import.run", "client.list")
    ]
    staff, _ = hire(owner, s.cmp, "staff@example.com", rights)
    assert staff.get(f"/api/c/{s.cmp}/imports/template/clients").status_code == 403
    content = workbook(template(owner, s.cmp, "clients"), "clients", [])
    assert upload(staff, s.cmp, "clients", content).status_code == 403
    wrong = upload(owner, s.cmp, "vehicles", content)
    assert wrong.json()["error"]["code"] == "IMPORT_TEMPLATE_UNSUPPORTED"
