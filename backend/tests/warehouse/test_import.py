"""XLSX import of the catalog (ТЗ § 22.2, § 22.4.1, § 22.5.1)."""

import io
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook

from c4rsys.warehouse import imports
from c4rsys.warehouse.models import Item
from tests.accounts.helpers import post
from tests.companies.helpers import create_company, hire, key

pytestmark = pytest.mark.django_db


@pytest.fixture
def cmp(owner: Client) -> str:
    public_id: str = create_company(owner, timezone="Asia/Bangkok").json()["company"]["public_id"]
    return public_id


def workbook(rows: list[dict[str, str]], extra_column: bool = False) -> bytes:
    book = load_workbook(io.BytesIO(imports.template()))
    sheet = book["items"]
    sheet.delete_rows(2)  # the example row
    header = [c.value for c in sheet[1]]
    if extra_column:
        sheet.cell(row=1, column=len(header) + 1, value="color")
    for row in rows:
        sheet.append([row.get(str(code), "") for code in header])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def upload(client: Client, cmp: str, content: bytes) -> Any:
    return client.post(
        f"/api/c/{cmp}/warehouse/imports",
        data={"file": SimpleUploadedFile("items.xlsx", content)},
        headers=key(),
    )


def test_template_has_codes_help_and_version(owner: Client, cmp: str) -> None:
    response = owner.get(f"/api/c/{cmp}/warehouse/imports/template")

    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert [c.value for c in book["items"][1]][:2] == ["name", "unit"]
    assert {"items", "help", "c4rsys"} <= set(book.sheetnames)


def test_check_shows_errors_and_matches_then_post_creates_rows(owner: Client, cmp: str) -> None:
    post(owner, f"/api/c/{cmp}/warehouse/items", {"name": "Wiper", "unit": "piece"}, key())
    content = workbook(
        [
            {"name": "Oil 5W-30", "unit": "liter", "manufacturer": "Shell", "barcode": "111"},
            {"name": "Wiper", "unit": "piece"},  # possible duplicate
            {"name": "", "unit": "piece"},  # error: no name
            {"name": "Filter", "unit": "box"},  # error: unknown unit
            {"name": "Pad", "unit": "шт", "side": "left", "condition": "new"},
        ],
        extra_column=True,
    )

    checked = upload(owner, cmp, content)

    assert checked.status_code == 201, checked.content
    operation = checked.json()["import"]
    actions = [r["action"] for r in operation["rows"]]
    assert actions == ["create", "skip", "error", "error", "create"]
    assert operation["rows"][1]["matches"][0]["by"] == "name_unit"
    assert {"code": "UNKNOWN_COLUMN", "column": "color"} in operation["warnings"]
    assert Item.objects.count() == 1  # nothing created by the check

    url = f"/api/c/{cmp}/warehouse/imports/{operation['id']}"
    decided = post(
        owner,
        f"{url}/decisions",
        {"decisions": [{"row": "R0002", "action": "create_duplicate"}], "expected_version": 1},
        key(),
    ).json()["import"]
    posted = post(owner, f"{url}/post", {"expected_version": decided["version"]}, key())

    assert posted.status_code == 200, posted.content
    body = posted.json()["import"]
    assert (body["loaded_count"], body["skipped_count"]) == (3, 2)
    assert Item.objects.filter(possible_duplicate=True).count() == 1
    oil = Item.objects.get(name="Oil 5W-30")
    assert oil.manufacturer_other == "Shell" and oil.barcodes.get().code == "111"
    assert Item.objects.get(name="Pad").unit == "piece"

    again = upload(owner, cmp, content).json()["import"]
    assert any(w["code"] == "REPEATED_FILE" for w in again["warnings"])


def test_errors_cannot_be_created_and_conflicts_are_not_posted(owner: Client, cmp: str) -> None:
    operation = upload(
        owner, cmp, workbook([{"name": "", "unit": "piece"}, {"name": "Lamp", "unit": "piece"}])
    ).json()["import"]
    url = f"/api/c/{cmp}/warehouse/imports/{operation['id']}"

    refused = post(
        owner,
        f"{url}/decisions",
        {"decisions": [{"row": "R0001", "action": "create"}], "expected_version": 1},
        key(),
    )
    assert refused.json()["error"]["code"] == "IMPORT_ROW_HAS_ERRORS"

    # Created by hand between the check and the posting: the row becomes a conflict (§ 22.2.2).
    post(owner, f"/api/c/{cmp}/warehouse/items", {"name": "Lamp", "unit": "piece"}, key())
    body = post(owner, f"{url}/post", {"expected_version": 1}, key()).json()["import"]

    assert body["rows"][1]["result"] == "conflict"
    assert Item.objects.filter(name="Lamp").count() == 1


def test_cancel_removes_created_cards_as_a_whole(owner: Client, cmp: str) -> None:
    operation = upload(owner, cmp, workbook([{"name": "Bulb", "unit": "piece"}])).json()["import"]
    url = f"/api/c/{cmp}/warehouse/imports/{operation['id']}"
    posted = post(owner, f"{url}/post", {"expected_version": 1}, key()).json()["import"]

    cancelled = post(owner, f"{url}/cancel", {"expected_version": posted["version"]}, key())

    assert cancelled.json()["import"]["state"] == "cancelled"
    assert not Item.objects.filter(name="Bulb").exists()


def test_other_files_are_refused_and_import_needs_its_right(owner: Client, cmp: str) -> None:
    assert upload(owner, cmp, b"not a workbook").json()["error"]["code"] == "IMPORT_FILE_NOT_XLSX"
    rights = [
        {
            "type": "grant",
            "permission": p,
            "locations": "all",
            "contexts": "all",
            "cash_accounts": "all",
        }
        for p in ("warehouse.view", "warehouse.catalog.manage")
    ]
    staff, _ = hire(owner, cmp, "staff@example.com", rights)
    assert upload(staff, cmp, workbook([{"name": "X", "unit": "piece"}])).status_code == 403
