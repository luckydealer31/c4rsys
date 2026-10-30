"""XLSX import: cancel of rows sharing one stock row, the history by the right of the section,
packaging confirmation with matches, lot prices by form, files after their keeping period and
the later history of initial stock (ТЗ § 22.2.2–22.2.5, § 22.4.2, § 22.5.1)."""

import io
from datetime import timedelta
from decimal import Decimal
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook

from c4rsys.files.models import StoredFile
from c4rsys.warehouse import imports
from c4rsys.warehouse.models import Asset, ImportOperation, Lot, LotStock, Packaging
from tests.accounts.helpers import post
from tests.companies.helpers import create_company, hire, key
from tests.rentals.scenario import build
from tests.warehouse.test_import_data import checked, run
from tests.warehouse.test_receipts import item, warehouse

pytestmark = pytest.mark.django_db


@pytest.fixture
def cmp(owner: Client) -> str:
    public_id: str = create_company(owner, timezone="Asia/Bangkok", main_currency="THB").json()[
        "company"
    ]["public_id"]
    return public_id


@pytest.fixture
def setup(owner: Client, cmp: str) -> dict[str, Any]:
    ctx = warehouse(owner, cmp)
    pads = item(owner, cmp, "Brake pads", "piece")
    oil = item(owner, cmp, "Engine oil", "liter")
    packed = post(
        owner,
        f"/api/c/{cmp}/warehouse/items/{oil['id']}/packagings",
        {"quantity": "4", "name": "Canister 4 l"},
        key(),
    )
    assert packed.status_code == 201, packed.content
    return {"ctx": ctx, "pads": pads["id"], "oil": oil["id"]}


def decide(owner: Client, cmp: str, operation: dict[str, Any], *pairs: tuple[str, str]) -> Any:
    return run(
        owner,
        cmp,
        operation,
        "decisions",
        decisions=[{"row": row, "action": action, "target": ""} for row, action in pairs],
    )


def test_cancel_takes_a_shared_stock_row_once(
    owner: Client, cmp: str, setup: dict[str, Any]
) -> None:
    a1 = {"item_id": setup["oil"], "context_id": setup["ctx"], "lot_code": "A1"}
    operation = checked(
        owner,
        cmp,
        "stock",
        [
            {**a1, "packaging_quantity": "4", "packs": "3"},
            {**a1, "packaging_quantity": "4", "packs": "2"},
        ],
    )
    posted = run(owner, cmp, operation, "post").json()["import"]
    assert [r["result"] for r in posted["rows"]] == ["created", "created"]
    row = LotStock.objects.get(lot=Lot.objects.get(), packaging=Packaging.objects.get())
    assert row.quantity == Decimal("20.00")

    # Part of the row is gone: each line alone would fit, their sum does not.
    LotStock.objects.filter(pk=row.pk).update(quantity=Decimal("15.00"))
    blocked = run(owner, cmp, posted, "cancel")
    assert blocked.json()["error"]["code"] == "IMPORT_CANCEL_BLOCKED"
    assert LotStock.objects.get(pk=row.pk).quantity == Decimal("15.00")

    LotStock.objects.filter(pk=row.pk).update(quantity=Decimal("20.00"))
    done = run(owner, cmp, posted, "cancel")
    assert done.status_code == 200, done.content
    assert LotStock.objects.get(pk=row.pk).quantity == Decimal("0.00")  # no lost update


def test_history_needs_the_right_of_the_section(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [{"kind": "person", "first_name": "Anna", "last_name": "Lee", "phone": "+6681111"}],
    )
    rights = [
        {
            "type": "grant",
            "permission": p,
            "locations": "all",
            "contexts": "all",
            "cash_accounts": "all",
        }
        for p in ("import.run", "import.history.view", "client.list")
    ]
    staff, _ = hire(owner, s.cmp, "history@example.com", rights)
    url = f"/api/c/{s.cmp}/imports"
    assert [o["id"] for o in owner.get(url).json()["imports"]] == [operation["id"]]
    assert staff.get(url).json()["imports"] == []
    assert staff.get(f"{url}?kind=clients").json()["imports"] == []
    assert staff.get(f"{url}/{operation['id']}").status_code == 403
    assert staff.get(f"{url}/{operation['id']}/file").status_code == 403
    assert owner.get(f"{url}/{operation['id']}/file").status_code == 200


def test_new_packaging_respects_matches(owner: Client, cmp: str, setup: dict[str, Any]) -> None:
    lot = {"item_id": setup["oil"], "context_id": setup["ctx"]}
    first = checked(
        owner,
        cmp,
        "stock",
        [{**lot, "lot_code": "D1", "lot_external_number": "X-1", "quantity": "1"}],
    )
    run(owner, cmp, first, "post")
    operation = checked(
        owner,
        cmp,
        "stock",
        [
            {
                **lot,
                "lot_code": "M1",
                "lot_external_number": "X-1",
                "packaging_quantity": "5",
                "packs": "1",
            }
        ],
    )
    row = operation["rows"][0]
    assert row["matches"] and row["extra"]["packaging_missing"] and row["action"] == "skip"
    # A possible duplicate is not created without the confirmation of its packaging.
    unconfirmed = decide(owner, cmp, operation, ("R0001", "create_duplicate"))
    assert unconfirmed.json()["error"]["code"] == "IMPORT_PACKAGING_CONFIRM"
    # The confirmation does not pass over the match: the row still needs its own choice.
    confirmed = decide(owner, cmp, operation, ("R0001", "create_packaging")).json()["import"]
    assert confirmed["rows"][0]["action"] == "skip"
    assert confirmed["rows"][0]["packaging_confirmed"] is True
    refused = decide(owner, cmp, confirmed, ("R0001", "create"))
    assert refused.json()["error"]["code"] == "IMPORT_ROW_HAS_MATCH"
    duplicate = decide(owner, cmp, confirmed, ("R0001", "create_duplicate")).json()["import"]
    assert duplicate["rows"][0]["action"] == "create_duplicate"
    posted = run(owner, cmp, duplicate, "post").json()["import"]
    assert posted["rows"][0]["result"] == "created", posted["rows"][0]
    new = Packaging.objects.get(quantity=5)
    assert Lot.objects.filter(external_number="X-1").count() == 2
    assert LotStock.objects.get(packaging=new).quantity == Decimal("5.00")


def test_price_of_a_new_packaging_is_not_compared_with_loose(
    owner: Client, cmp: str, setup: dict[str, Any]
) -> None:
    oil = setup["oil"]
    mode = post(owner, f"/api/c/{cmp}/warehouse/items/{oil}/price-mode", {"mode": "lot"}, key())
    assert mode.status_code == 200, mode.content
    lot = {"item_id": oil, "context_id": setup["ctx"]}
    first = checked(
        owner, cmp, "stock", [{**lot, "lot_code": "P1", "quantity": "3", "sale_price": "100"}]
    )
    assert run(owner, cmp, first, "post").json()["import"]["rows"][0]["result"] == "created"
    existing = Lot.objects.get()
    more = checked(
        owner,
        cmp,
        "stock",
        [
            {
                **lot,
                "lot_id": existing.public_id,
                "packaging_quantity": "5",
                "packs": "1",
                "sale_price": "450",
            },
            {**lot, "lot_id": existing.public_id, "quantity": "1", "sale_price": "120"},
        ],
    )
    assert more["rows"][0]["errors"] == []  # a packaged price against the loose one
    loose = {"field": "sale_price", "code": "LOT_PRICE_EXISTS"}
    assert more["rows"][1]["errors"] == [loose]  # the loose price is compared with the loose


def test_corrected_file_after_removal_is_not_kept(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [
            {"kind": "person", "first_name": "Anna", "last_name": "Lee", "phone": "+6681111"},
            {"kind": "person", "first_name": "Bob", "last_name": "", "phone": "+6682222"},
        ],
    )
    posted = run(owner, s.cmp, operation, "post").json()["import"]
    url = f"/api/c/{s.cmp}/imports/{operation['id']}"
    correction = owner.get(f"{url}/correction")
    created = ImportOperation.objects.get(pk=operation["id"]).created_at
    assert imports.remove_old_files(imports.shift_months(created, 6)) == 1

    book = load_workbook(io.BytesIO(correction.content))
    sheet = book["clients"]
    header = [c.value for c in sheet[1]]
    sheet.delete_rows(2, 1)
    values = {"row_code": "R0002", "kind": "person", "first_name": "Bob", "last_name": "Ray"}
    sheet.append([{**values, "phone": "+6682222"}.get(str(c), "") for c in header])
    output = io.BytesIO()
    book.save(output)
    kept = StoredFile.objects.filter(purpose=StoredFile.Purpose.IMPORT_FILE).count()

    bad = owner.post(
        f"{url}/reupload",
        data={"expected_version": "x1", "file": SimpleUploadedFile("f.xlsx", output.getvalue())},
        headers=key(),
    )
    assert bad.status_code == 400 and bad.json()["error"]["field"] == "expected_version"

    reuploaded = owner.post(
        f"{url}/reupload",
        data={
            "expected_version": str(posted["version"]),
            "file": SimpleUploadedFile("fixed.xlsx", output.getvalue()),
        },
        headers=key(),
    )
    assert reuploaded.status_code == 200, reuploaded.content
    # The keeping period ended with the first upload: the later file is not stored (§ 22.2.4).
    assert StoredFile.objects.filter(purpose=StoredFile.Purpose.IMPORT_FILE).count() == kept
    operation_now = ImportOperation.objects.get(pk=operation["id"])
    assert operation_now.uploads[-1]["file"] is None
    assert imports.remove_old_files(imports.shift_months(created, 7)) == 0


def test_corrected_file_before_removal_is_removed_with_the_first(owner: Client) -> None:
    s = build(owner)
    operation = checked(
        owner,
        s.cmp,
        "clients",
        [{"kind": "person", "first_name": "Bob", "last_name": "", "phone": "+6682222"}],
    )
    posted = run(owner, s.cmp, operation, "post").json()["import"]
    url = f"/api/c/{s.cmp}/imports/{operation['id']}"
    book = load_workbook(io.BytesIO(owner.get(f"{url}/correction").content))
    output = io.BytesIO()
    book.save(output)
    reuploaded = owner.post(
        f"{url}/reupload",
        data={
            "expected_version": str(posted["version"]),
            "file": SimpleUploadedFile("again.xlsx", output.getvalue()),
        },
        headers=key(),
    )
    assert reuploaded.status_code == 200, reuploaded.content
    created = ImportOperation.objects.get(pk=operation["id"]).created_at
    assert imports.remove_old_files(created + timedelta(days=150)) == 0
    assert imports.remove_old_files(imports.shift_months(created, 6)) == 1
    found = ImportOperation.objects.get(pk=operation["id"])
    ids = [found.file_id, *[u["file"] for u in found.uploads]]
    assert StoredFile.objects.filter(pk__in=ids, removed_at__isnull=True).count() == 0


def test_later_history_of_initial_stock_blocks_cancel(
    owner: Client, cmp: str, setup: dict[str, Any]
) -> None:
    operation = checked(
        owner, cmp, "stock", [{"item_id": setup["pads"], "context_id": setup["ctx"]}]
    )
    posted = run(owner, cmp, operation, "post").json()["import"]
    asset = Asset.objects.get()
    priced = owner.put(
        f"/api/c/{cmp}/warehouse/assets/{asset.id}/price",
        data={"amount": "750"},
        content_type="application/json",
        headers=key(),
    )
    assert priced.status_code == 200, priced.content
    refused = run(owner, cmp, posted, "cancel")
    assert refused.json()["error"]["code"] == "IMPORT_CANCEL_BLOCKED"
    blocker = refused.json()["error"]["params"]["blockers"][0]
    assert blocker == {
        "reason": "history",
        "id": str(asset.id),
        "action": "warehouse.asset_price_set",
    }
    asset.refresh_from_db()
    assert asset.state == "in_stock"
