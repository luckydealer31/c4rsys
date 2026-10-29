"""XLSX import (ТЗ § 22.2–22.5): one mechanism for the templates of the catalog, vehicles,
clients and counterparties.

Two steps: the file is checked and its errors and matches are shown; the user chooses what to do
with each row and confirms. Correct rows are loaded, wrong ones skipped with their reasons. A
match never merges or overwrites an existing record: the row is skipped, linked to the record or
created as a possible duplicate. Just before posting the rows are checked again against current
data. The operation keeps its file, author, time, counts and links; it is cancelled only as a
whole while its records are unused and have no history of their own (§ 22.2.5).

A corrected file is uploaded from the history of the import: the file for correction carries the
codes of the rows not loaded, rows are matched by «import + row code», loaded rows are never
posted again (§ 22.2.7). Source files are removed six months after the first upload (§ 22.2.4).

The kinds of data are adapters of their modules: each checks its rows and creates or removes its
own records through its own services. The operation itself is kept by the warehouse module,
which had the first import.
"""

import hashlib
import io
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Protocol
from uuid import UUID
from zipfile import BadZipFile

from django.db import transaction
from django.db.models import ProtectedError
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from c4rsys.companies.models import Company
from c4rsys.files import services as files
from c4rsys.files.models import StoredFile
from c4rsys.platform import errors
from c4rsys.platform.commands import Actor, CommandContext
from c4rsys.platform.models import AuditEvent
from c4rsys.warehouse.models import ImportOperation

META_SHEET = "c4rsys"
HELP_SHEET = "help"
PROBLEMS_SHEET = "problems"
ROW_CODE = "row_code"
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
KEEP_FILES_MONTHS = 6
LOADED = ("created", "linked")
ACTIONS = {"create", "skip", "link", "create_duplicate"}


@dataclass
class Scope:
    """What a row is checked against: the company, the employee (for the area of access,
    § 22.5.1) and values already met in the file."""

    company_id: UUID
    actor: Actor
    seen: dict[str, set[str]] = field(default_factory=dict)

    def repeated(self, key: str, value: str) -> bool:
        """True when the value was already met in the file; remembers it otherwise."""
        bucket = self.seen.setdefault(key, set())
        if value in bucket:
            return True
        bucket.add(value)
        return False


@dataclass
class Checked:
    errors: list[dict[str, str]] = field(default_factory=list)
    # Existing records: {"id", "name", "by"}; several are shown without a choice (§ 22.2.3).
    matches: list[dict[str, str]] = field(default_factory=list)
    # The ID points to one record and other identifiers to another: resolved before posting by
    # linking or skipping (§ 22.2.3).
    conflict: bool = False
    # Values the module derived from the row, kept with it for posting.
    extra: dict[str, Any] = field(default_factory=dict)

    def error(self, field_code: str, code: str) -> None:
        self.errors.append({"field": field_code, "code": code})

    def match(self, record_id: UUID | str, name: str, by: str) -> None:
        if all(m["id"] != str(record_id) for m in self.matches):
            self.matches.append({"id": str(record_id), "name": name, "by": by})


class Adapter(Protocol):
    kind: str
    version: int
    sheet: str
    # Technical codes of the columns (§ 22.2.6) with their maximum length.
    columns: dict[str, int]
    required: tuple[str, ...]
    # code -> (English, Russian, example) for the help sheet.
    help: dict[str, tuple[str, str, str]]
    # The usual right to the section the data belongs to (§ 22.5.1).
    permission: str
    # Journal object type of the created records, for their later history (§ 22.2.5).
    object_type: str

    def references(self, company_id: UUID) -> list[tuple[str, list[list[str]]]]:
        """Extra sheets of the template, such as the IDs of the company's locations."""
        ...

    def check_row(self, scope: Scope, values: dict[str, str]) -> Checked: ...

    def create(
        self,
        ctx: CommandContext,
        company: Company,
        values: dict[str, str],
        extra: dict[str, Any],
        *,
        duplicate: bool,
        operation_id: UUID,
    ) -> UUID: ...

    def remove(self, ctx: CommandContext, company_id: UUID, record_ids: list[UUID]) -> None:
        """Deletes the records created by the import; a reference of another record raises
        ProtectedError, a confirmed link raises IMPORT_CANCEL_BLOCKED."""
        ...


KINDS = ("items", "vehicles", "clients", "counterparties")


def adapter(kind: str) -> Adapter:
    if kind == "items":
        from c4rsys.warehouse.import_items import ADAPTER as items

        return items
    if kind == "vehicles":
        from c4rsys.fleet.imports import ADAPTER as vehicles

        return vehicles
    if kind == "clients":
        from c4rsys.parties.imports import CLIENTS

        return CLIENTS
    if kind == "counterparties":
        from c4rsys.parties.imports import COUNTERPARTIES

        return COUNTERPARTIES
    raise errors.validation_error("kind")


# --- Template and reading of files (§ 22.2.6) --------------------------------------------------


def _book(kind: Adapter, company_id: UUID | None) -> Workbook:
    book = Workbook()
    sheet = book.create_sheet(kind.sheet, 0)
    for extra in book.sheetnames[1:]:
        del book[extra]
    sheet.append(list(kind.columns))
    help_sheet = book.create_sheet(HELP_SHEET)
    help_sheet.append(["code", "required", "en", "ru", "example"])
    for code in kind.columns:
        en, ru, example = kind.help[code]
        help_sheet.append([code, "yes" if code in kind.required else "", en, ru, example])
    if company_id is not None:
        for title, rows in kind.references(company_id):
            reference = book.create_sheet(title)
            for row in rows:
                reference.append(row)
    meta = book.create_sheet(META_SHEET)
    meta.append(["template", kind.kind])
    meta.append(["version", kind.version])
    meta.sheet_state = "hidden"
    return book


def _bytes(book: Workbook) -> bytes:
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def template(kind: str = "items", company_id: UUID | None = None) -> bytes:
    """Technical codes in the first row and an example in the second, explanations on the help
    sheet, the template and its version on the hidden service sheet."""
    found = adapter(kind)
    book = _book(found, company_id)
    book[found.sheet].append([found.help[c][2] for c in found.columns])
    help_sheet = book[HELP_SHEET]
    help_sheet.append([])
    help_sheet.append(
        [
            "",
            "",
            f"The second row of the «{found.sheet}» sheet is an example: replace or delete it.",
            f"Вторая строка листа «{found.sheet}» — пример: замените или удалите её.",
            "",
        ]
    )
    return _bytes(book)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass(frozen=True)
class Parsed:
    rows: list[dict[str, Any]]
    warnings: list[dict[str, Any]]


def parse(content: bytes, kind: str = "items", *, correction: bool = False) -> Parsed:
    """Rows of the data sheet as text values; the file must be the template of this kind and
    version (§ 22.2.6). A corrected file must have the codes of its rows (§ 22.2.7)."""
    found = adapter(kind)
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, ValueError, OSError) as exc:
        raise errors.validation_error("file", code="IMPORT_FILE_NOT_XLSX") from exc
    if META_SHEET not in book.sheetnames or found.sheet not in book.sheetnames:
        raise errors.validation_error("file", code="IMPORT_TEMPLATE_UNSUPPORTED")
    meta = {
        _cell(r[0]): _cell(r[1])
        for r in book[META_SHEET].iter_rows(values_only=True)
        if r and len(r) > 1
    }
    if meta.get("template") != found.kind or meta.get("version") != str(found.version):
        raise errors.validation_error("file", code="IMPORT_TEMPLATE_UNSUPPORTED")
    rows = list(book[found.sheet].iter_rows(values_only=True))
    if not rows:
        raise errors.validation_error("file", code="IMPORT_TEMPLATE_UNSUPPORTED")
    header = [_cell(c) for c in rows[0]]
    warnings = [
        {"code": "UNKNOWN_COLUMN", "column": c}
        for c in header
        if c and c not in found.columns and c != ROW_CODE
    ]
    missing = [c for c in found.required if c not in header]
    if correction and ROW_CODE not in header:
        missing.append(ROW_CODE)
    if missing:
        raise errors.validation_error("file", code="IMPORT_COLUMNS_MISSING", columns=missing)
    result = []
    for number, row in enumerate(rows[1:], start=2):
        cells = {code: _cell(row[i]) if i < len(row) else "" for i, code in enumerate(header)}
        values = {code: value for code, value in cells.items() if code in found.columns}
        if not any(values.values()):
            continue  # empty rows are ignored (§ 22.2.6)
        result.append({"number": number, "values": values, "row_code": cells.get(ROW_CODE, "")})
    if len(result) > MAX_ROWS:
        raise errors.validation_error("file", code="IMPORT_TOO_MANY_ROWS", max_rows=MAX_ROWS)
    return Parsed(result, warnings)


def _checked(kind: Adapter, scope: Scope, values: dict[str, str]) -> Checked:
    result = kind.check_row(scope, values)
    for code, length in kind.columns.items():
        if len(values.get(code, "")) > length:
            result.error(code, "FIELD_TOO_LONG")
    return result


def _default_action(row: dict[str, Any]) -> str:
    if row["errors"]:
        return "error"
    return "skip" if row["matches"] or row["conflict"] else "create"


def _apply_check(row: dict[str, Any], checked: Checked) -> None:
    row["errors"] = checked.errors
    row["matches"] = checked.matches
    row["conflict"] = checked.conflict
    row["extra"] = checked.extra
    row["action"] = _default_action(row)
    row["target"] = None
    row["result"] = None
    row["record_id"] = None


# --- Check, decisions, posting ------------------------------------------------------------------


def _store(ctx: CommandContext, company: Company, content: bytes) -> StoredFile:
    stored = files.store(
        purpose=StoredFile.Purpose.IMPORT_FILE,
        uploaded_by=ctx.actor.user_id,
        content=content,
        content_type=XLSX,
    )
    files.attach_to_company(stored.id, company.id)
    return stored


def check(
    ctx: CommandContext, company: Company, content: bytes, file_name: str, kind: str = "items"
) -> ImportOperation:
    """Step one: the file is stored with the operation, rows are checked; nothing is created."""
    found = adapter(kind)
    if len(content) > MAX_FILE_BYTES:
        raise errors.validation_error("file", code="FILE_TOO_LARGE", max_bytes=MAX_FILE_BYTES)
    parsed = parse(content, kind)
    digest = hashlib.sha256(content).hexdigest()
    warnings = list(parsed.warnings)
    previous = (
        ImportOperation.objects.filter(company=company, kind=kind, sha256=digest)
        .exclude(state=ImportOperation.State.CANCELLED)
        .order_by("-created_at")
        .first()
    )
    if previous is not None:
        # § 22.2.7: the same file again — a warning with the previous result.
        warnings.append(
            {
                "code": "REPEATED_FILE",
                "import": str(previous.id),
                "loaded": previous.loaded_count,
                "skipped": previous.skipped_count,
            }
        )
    stored = _store(ctx, company, content)
    scope = Scope(company.id, ctx.actor)
    rows = []
    for index, row in enumerate(parsed.rows, start=1):
        entry: dict[str, Any] = {
            "number": row["number"],
            "values": row["values"],
            "code": f"R{index:04d}",
        }
        _apply_check(entry, _checked(found, scope, row["values"]))
        rows.append(entry)
    operation = ImportOperation.objects.create(
        company=company,
        kind=kind,
        template_version=found.version,
        file=stored,
        file_name=file_name[:255],
        sha256=digest,
        rows=rows,
        warnings=warnings,
        created_by_user_id=ctx.actor.user_id,
    )
    ctx.audit(
        "warehouse.import_checked",
        object_type="warehouse_import",
        object_id=operation.id,
        after={
            "kind": kind,
            "file": file_name[:255],
            "rows": len(rows),
            "errors": sum(1 for r in rows if r["errors"]),
            "matches": sum(1 for r in rows if r["matches"]),
        },
    )
    return operation


def _lock(
    company_id: UUID, operation_id: UUID, expected_version: int | None = None
) -> ImportOperation:
    operation = (
        ImportOperation.objects.select_for_update()
        .filter(company_id=company_id, pk=operation_id)
        .first()
    )
    if operation is None:
        raise errors.not_found()
    if expected_version is not None and operation.version != expected_version:
        raise errors.version_conflict()
    return operation


def decide(
    ctx: CommandContext,
    company_id: UUID,
    operation_id: UUID,
    decisions: list[tuple[str, str, str | None]],
    expected_version: int,
) -> ImportOperation:
    """The user's choice for rows: create, skip, link to an existing record or create a possible
    duplicate (§ 22.2.3). A row with errors can only be skipped; a row whose ID and other
    identifiers point to different records is linked to one of them or skipped."""
    operation = _lock(company_id, operation_id, expected_version)
    if operation.state != ImportOperation.State.CHECKED:
        raise errors.business_rule("IMPORT_NOT_CHECKED")
    by_code = {r["code"]: r for r in operation.rows}
    for code, action, target in decisions:
        row = by_code.get(code)
        if row is None:
            raise errors.validation_error("decisions", code="IMPORT_ROW_UNKNOWN", row=code)
        if action not in ACTIONS:
            raise errors.validation_error("decisions", code="IMPORT_ACTION_UNKNOWN", row=code)
        if row.get("result") in LOADED:
            raise errors.business_rule("IMPORT_ROW_POSTED", row=code)
        if row["errors"] and action != "skip":
            raise errors.business_rule("IMPORT_ROW_HAS_ERRORS", row=code)
        if action == "link" and target not in {m["id"] for m in row["matches"]}:
            raise errors.validation_error("decisions", code="IMPORT_LINK_TARGET", row=code)
        if action == "create" and row["matches"]:
            raise errors.business_rule("IMPORT_ROW_HAS_MATCH", row=code)
        if action in ("create", "create_duplicate") and row.get("conflict"):
            raise errors.business_rule("IMPORT_ROW_ID_CONFLICT", row=code)
        row["action"] = action
        row["target"] = target if action == "link" else None
    operation.version += 1
    operation.save(update_fields=["rows", "version"])
    ctx.audit(
        "warehouse.import_decided",
        object_type="warehouse_import",
        object_id=operation.id,
        after={"decisions": [{"row": c, "action": a, "target": t} for c, a, t in decisions]},
    )
    return operation


def _problem(exc: errors.ApiError) -> dict[str, str]:
    return {"field": exc.field or "", "code": exc.code}


def post(
    ctx: CommandContext, company: Company, operation_id: UUID, expected_version: int
) -> ImportOperation:
    """Step two: rows are checked again against current data; a changed result is a conflict
    and the row is not posted (§ 22.2.2). Each row is posted on its own; rows loaded before stay
    as they are (§ 22.2.7)."""
    operation = _lock(company.id, operation_id, expected_version)
    if operation.state != ImportOperation.State.CHECKED:
        raise errors.business_rule("IMPORT_NOT_CHECKED")
    found = adapter(operation.kind)
    scope = Scope(company.id, ctx.actor)
    rows = []
    for row in operation.rows:
        if row.get("result") in LOADED:
            rows.append(row)
            continue
        current = _checked(found, scope, row["values"])
        action = row["action"]
        result: dict[str, Any] = {**row, "errors": current.errors, "extra": current.extra}
        known = {m["id"] for m in row["matches"]}
        fresh = {m["id"] for m in current.matches}
        if action in ("skip", "error"):
            result["result"] = "skipped"
        elif current.errors or (
            action in ("create", "create_duplicate") and (fresh - known or current.conflict)
        ):
            result["result"] = "conflict"
            result["matches"] = current.matches
            result["conflict"] = current.conflict
        elif action == "link":
            if row["target"] not in fresh:
                result["result"] = "conflict"
                result["matches"] = current.matches
            else:
                result["result"] = "linked"
                result["record_id"] = row["target"]
                result["request"] = ctx.request_id
        else:
            mark = ctx.audit_mark()
            try:
                with transaction.atomic():
                    record_id = found.create(
                        ctx,
                        company,
                        row["values"],
                        current.extra,
                        duplicate=action == "create_duplicate",
                        operation_id=operation.id,
                    )
            except errors.ApiError as exc:
                ctx.discard_audit(mark)
                result["result"] = "conflict"
                result["errors"] = [*current.errors, _problem(exc)]
            else:
                result["result"] = "created"
                result["record_id"] = str(record_id)
                result["request"] = ctx.request_id
        rows.append(result)
    operation.rows = rows
    operation.state = ImportOperation.State.POSTED
    operation.posted_at = timezone.now()
    operation.loaded_count = sum(1 for r in rows if r.get("result") in LOADED)
    operation.skipped_count = len(rows) - operation.loaded_count
    operation.version += 1
    operation.save()
    ctx.audit(
        "warehouse.import_posted",
        object_type="warehouse_import",
        object_id=operation.id,
        after={"loaded": operation.loaded_count, "skipped": operation.skipped_count},
    )
    return operation


# --- Corrected file (§ 22.2.7) -----------------------------------------------------------------


def correction(operation: ImportOperation) -> bytes:
    """The rows not loaded with their codes and values, and the reasons on a separate sheet."""
    if operation.state != ImportOperation.State.POSTED:
        raise errors.business_rule("IMPORT_NOT_POSTED")
    found = adapter(operation.kind)
    book = _book(found, operation.company_id)
    sheet = book[found.sheet]
    sheet.delete_rows(1)
    columns = list(found.columns)
    sheet.append([ROW_CODE, *columns])
    problems = book.create_sheet(PROBLEMS_SHEET, 1)
    problems.append([ROW_CODE, "number", "result", "field", "code"])
    for row in operation.rows:
        if row.get("result") in LOADED:
            continue
        sheet.append([row["code"], *[row["values"].get(c, "") for c in columns]])
        reasons = row["errors"] or [{"field": "", "code": "MATCH" if row["matches"] else "SKIPPED"}]
        for reason in reasons:
            problems.append(
                [
                    row["code"],
                    row["number"],
                    row.get("result") or "",
                    reason["field"],
                    reason["code"],
                ]
            )
    return _bytes(book)


def reupload(
    ctx: CommandContext,
    company: Company,
    operation_id: UUID,
    content: bytes,
    file_name: str,
    expected_version: int,
) -> ImportOperation:
    """A corrected file within the same import: rows are matched by their codes, the order may
    change. A missing or repeated code is an error of that row and nothing is guessed; loaded
    rows are not posted again even when their values changed. The rows of the file are checked
    anew and wait for the user's decision and posting (§ 22.2.7)."""
    operation = _lock(company.id, operation_id, expected_version)
    if operation.state != ImportOperation.State.POSTED:
        raise errors.business_rule("IMPORT_NOT_POSTED")
    if len(content) > MAX_FILE_BYTES:
        raise errors.validation_error("file", code="FILE_TOO_LARGE", max_bytes=MAX_FILE_BYTES)
    found = adapter(operation.kind)
    parsed = parse(content, operation.kind, correction=True)
    stored = _store(ctx, company, content)
    by_code = {r["code"]: r for r in operation.rows}
    counts = Counter(r["row_code"] for r in parsed.rows if r["row_code"])
    warnings: list[dict[str, Any]] = list(parsed.warnings)
    scope = Scope(company.id, ctx.actor)
    corrected: set[str] = set()
    for entry in parsed.rows:
        code = entry["row_code"]
        problem = None
        if not code:
            problem = "ROW_CODE_MISSING"
        elif counts[code] > 1:
            problem = "ROW_CODE_REPEATED"
        elif code not in by_code:
            problem = "ROW_CODE_UNKNOWN"
        elif by_code[code].get("result") in LOADED:
            problem = "ROW_ALREADY_POSTED"
        if problem is not None:
            warnings.append({"code": problem, "number": entry["number"], "row": code})
            continue
        row = by_code[code]
        row["values"] = entry["values"]
        row["number"] = entry["number"]
        _apply_check(row, _checked(found, scope, entry["values"]))
        corrected.add(code)
    for row in operation.rows:
        if row["code"] not in corrected and row.get("result") not in LOADED:
            row["action"] = "skip"
    warnings.append({"code": "CORRECTED_FILE", "rows": len(corrected)})
    operation.uploads = [
        *operation.uploads,
        {
            "file": str(stored.id),
            "file_name": file_name[:255],
            "sha256": hashlib.sha256(content).hexdigest(),
            "at": timezone.now().isoformat(),
            "by": str(ctx.actor.user_id),
        },
    ]
    operation.warnings = warnings
    operation.state = ImportOperation.State.CHECKED
    operation.version += 1
    operation.save(update_fields=["rows", "uploads", "warnings", "state", "version"])
    ctx.audit(
        "warehouse.import_reuploaded",
        object_type="warehouse_import",
        object_id=operation.id,
        after={"file": file_name[:255], "rows": sorted(corrected)},
    )
    return operation


# --- Cancel (§ 22.2.5) --------------------------------------------------------------------------


def _blocker(obj: Any) -> dict[str, str]:
    return {"reason": "in_use", "type": obj._meta.db_table, "id": str(obj.pk)}


def cancel(
    ctx: CommandContext, company_id: UUID, operation_id: UUID, expected_version: int
) -> ImportOperation:
    """Only as a whole and only while the created records are unused and have no history after
    the import (§ 22.2.5); links to existing records are simply forgotten. The obstacles are
    shown; nothing is removed in part."""
    operation = _lock(company_id, operation_id, expected_version)
    if operation.state != ImportOperation.State.POSTED:
        raise errors.business_rule("IMPORT_NOT_POSTED")
    found = adapter(operation.kind)
    created = [r for r in operation.rows if r.get("result") == "created"]
    ids = [UUID(str(r.get("record_id") or r.get("item_id"))) for r in created]
    requests = {r["request"] for r in created if r.get("request")}
    checked_ids = [str(r["record_id"]) for r in created if r.get("request")]
    history = (
        AuditEvent.objects.filter(
            company_id=company_id, object_type=found.object_type, object_id__in=checked_ids
        )
        .exclude(request_id__in=requests)
        .order_by("occurred_at")[:20]
    )
    blockers = [{"reason": "history", "id": e.object_id, "action": e.action} for e in history]
    if blockers:
        raise errors.business_rule("IMPORT_CANCEL_BLOCKED", blockers=blockers)
    try:
        with transaction.atomic():
            found.remove(ctx, company_id, ids)
    except ProtectedError as exc:
        raise errors.business_rule(
            "IMPORT_CANCEL_BLOCKED",
            blockers=[_blocker(o) for o in list(exc.protected_objects)[:20]],
        ) from exc
    operation.state = ImportOperation.State.CANCELLED
    operation.cancelled_at = timezone.now()
    operation.version += 1
    operation.save(update_fields=["state", "cancelled_at", "version"])
    ctx.audit(
        "warehouse.import_cancelled",
        object_type="warehouse_import",
        object_id=operation.id,
        after={"kind": operation.kind, "deleted": [str(i) for i in ids]},
    )
    return operation


# --- Source files (§ 22.2.4) --------------------------------------------------------------------


def shift_months(moment: datetime, months: int) -> datetime:
    """The same day and time `months` calendar months later (earlier when negative); a missing
    day is the last of the month."""
    month = moment.month + months
    year = moment.year + (month - 1) // 12
    month = (month - 1) % 12 + 1
    day = moment.day
    while True:
        try:
            return moment.replace(year=year, month=month, day=day)
        except ValueError:
            day -= 1


def remove_old_files(now: datetime | None = None) -> int:
    """The source and corrected files of imports first uploaded six months ago or earlier; the
    operation and its rows stay (§ 22.2.4). Journaled as a system process."""
    moment = now or timezone.now()
    border = shift_months(moment, -KEEP_FILES_MONTHS)
    removed = 0
    for operation in ImportOperation.objects.filter(
        created_at__lte=border, files_removed_at__isnull=True
    ):
        with transaction.atomic():
            ids = [operation.file_id] if operation.file_id else []
            ids += [UUID(u["file"]) for u in operation.uploads]
            for file_id in ids:
                files.remove(file_id)
            operation.files_removed_at = moment
            operation.save(update_fields=["files_removed_at"])
            AuditEvent.objects.create(
                company_id=operation.company_id,
                system_process="import_files_retention",
                action="warehouse.import_files_removed",
                object_type="warehouse_import",
                object_id=str(operation.id),
                after={"files": [str(i) for i in ids]},
            )
        removed += 1
    return removed


def payload(operation: ImportOperation, with_rows: bool = True) -> dict[str, Any]:
    body: dict[str, Any] = {
        "id": str(operation.id),
        "version": operation.version,
        "kind": operation.kind,
        "state": operation.state,
        "file_name": operation.file_name,
        "created_at": operation.created_at.isoformat(),
        "posted_at": operation.posted_at.isoformat() if operation.posted_at else None,
        "cancelled_at": operation.cancelled_at.isoformat() if operation.cancelled_at else None,
        "loaded_count": operation.loaded_count,
        "skipped_count": operation.skipped_count,
        "rows_count": len(operation.rows),
        "warnings": operation.warnings,
        "uploads": [{"file_name": u["file_name"], "at": u["at"]} for u in operation.uploads],
        # Files are kept six months from the first upload; a corrected file does not extend it.
        "files_kept_until": shift_months(operation.created_at, KEEP_FILES_MONTHS).isoformat(),
        "files_removed": operation.files_removed_at is not None,
    }
    if with_rows:
        body["rows"] = [
            {**r, "record_id": r.get("record_id") or r.get("item_id")} for r in operation.rows
        ]
    return body
